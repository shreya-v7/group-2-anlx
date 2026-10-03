from pathlib import Path
import json
from html import escape
from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent; O=R/'documents'; O.mkdir(exist_ok=True)
pages=json.loads((R/'report_content.json').read_text())
S=getSampleStyleSheet()
S.add(ParagraphStyle(name='BodyX',fontName='Helvetica',fontSize=10.2,leading=14.4,spaceAfter=9))
S.add(ParagraphStyle(name='TitleX',fontName='Helvetica-Bold',fontSize=21,leading=25,spaceAfter=15))
S.add(ParagraphStyle(keepWithNext=True,name='SecX',fontName='Helvetica-Bold',fontSize=15.5,leading=20,spaceBefore=4,spaceAfter=8,textColor=colors.HexColor('#173e50')))
S.add(ParagraphStyle(keepWithNext=True,name='HeadX',fontName='Helvetica-Bold',fontSize=12,leading=16,spaceBefore=6,spaceAfter=7))
S.add(ParagraphStyle(name='CellX',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=0))
def par(text,style='BodyX'):return Paragraph(escape(text).replace('\n','<br/>'),S[style])
def tbl(data):
 widths=([104,407] if len(data[0])==2 else [511/len(data[0])]*len(data[0]))
 z=Table([[par(str(c),'CellX') for c in row] for row in data],colWidths=widths,repeatRows=1)
 z.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7eff2')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#b8c5cb'))]));return z
def footer(c,d):
 c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#60727a'));c.drawString(42,29,'CMU 95820 | Job posting extraction | October 2026');c.drawRightString(A4[0]-42,29,str(d.page))
def pdf(name,content):
 st=[]
 for i,pg in enumerate(content):
  if i and pg['title'].startswith('References'):st.append(PageBreak())
  elif i:st.append(Spacer(1,10))
  st.append(par(pg['title'],'TitleX' if i==0 else 'SecX'))
  for b in pg['blocks']:
   if 'text'in b:st.append(par(b['text']))
   elif 'heading'in b:st.append(par(b['heading'],'HeadX'))
   else:st.extend([tbl(b['table']),Spacer(1,12)])
 SimpleDocTemplate(str(O/name),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=38,bottomMargin=50,title=content[0]['title'],author='Shreya Verma; Chih-yu Li; William Sun').build(st,onFirstPage=footer,onLaterPages=footer)
pdf('Team_Research_Report.pdf',pages)
d=Document(); sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.65)
d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
for elem in list(d.styles.element.iter(qn('w:pBdr'))):elem.getparent().remove(elem)
for sn in ['Title','Heading 1','Heading 2']:
 d.styles[sn].font.color.rgb=RGBColor(0,0,0)
for i,pg in enumerate(pages):
 if i:d.add_page_break()
 d.add_heading(pg['title'],0 if i==0 else 1)
 for b in pg['blocks']:
  if 'text'in b:d.add_paragraph(b['text'])
  elif 'heading'in b:d.add_heading(b['heading'],2)
  else:
   tab=d.add_table(rows=0,cols=len(b['table'][0]));tab.style='Light Shading Accent 1'
   borders=__import__('docx.oxml',fromlist=['OxmlElement']).OxmlElement('w:tblBorders')
   for edge in ['top','left','bottom','right','insideH','insideV']:
    e=__import__('docx.oxml',fromlist=['OxmlElement']).OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tab._tbl.tblPr.append(borders)
   for ri,row in enumerate(b['table']):
    cells=tab.add_row().cells
    for j,v in enumerate(row):
     cells[j].text=str(v)
     for para in cells[j].paragraphs:
      for run in para.runs:run.font.size=Pt(9);run.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
d.save(O/'Team_Research_Report.docx')
md=[]
for pg in pages:
 md.append('# '+pg['title'])
 for b in pg['blocks']:
  if 'text'in b:md.append(b['text'])
  elif 'heading'in b:md.append('## '+b['heading'])
  else:md.append('\n'.join('| '+' | '.join(str(v) for v in row)+' |' for row in [b['table'][0],['---']*len(b['table'][0])]+b['table'][1:]))
(O/'Team_Research_Report.md').write_text('\n\n'.join(md))

